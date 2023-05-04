from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook
import os


nome_arquivo = "Alunos.xlsx"
planilhaAlunos = load_workbook(nome_arquivo)
sheet = planilhaAlunos['Nomes']

for L in range(2, len(sheet['A']) + 1):
    arquivo = Document('Certificado1.docx')

    estilo = arquivo.styles['Normal']

    nomeAluno = sheet['A%s' % L].value

    for p in arquivo.paragraphs:
        if "@nome" in p.text:
            p.text = nomeAluno
            fonte = estilo.font
            fonte.name = 'Calibri (Corpo)'
            fonte.size = Pt(24)

    certificados = f'C:\\Users\\Win\\Documents\\GitHub\\Python\\Word\\Certificados\\{nomeAluno}.docx'
    arquivo.save(certificados)
